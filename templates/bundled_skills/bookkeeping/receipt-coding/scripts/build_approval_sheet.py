#!/usr/bin/env python3
"""Approval sheet renderer, scheme reporter, and row validator.

Two modes.

  Report the active coding scheme (counts come from the file, never estimated):

      python3 build_approval_sheet.py --scheme-info

  Validate and render a coded batch into ONE approval sheet:

      python3 build_approval_sheet.py --rows-file /tmp/receipt_rows.json \\
                                      --out /tmp/approval_sheet

WHY THIS FILE DOES THE ENFORCING
--------------------------------
Rules written into SKILL.md asking the agent to count accurately, or to keep
confidence consistent with routing, did not hold across three runs. Rules
implemented here have held every time. So the consistency guarantees live in
code:

  * an account not present in the scheme cannot reach the sheet
  * an account marked auto_assign:false cannot be assigned from a receipt
  * a row routed to a fallback cannot also claim high account confidence
  * a row routed to a fallback is always flagged for review

Corrections are applied, not merely warned about, and every correction is
reported on stderr so the caller can describe them accurately.

``--rows-file`` exists so callers never need a heredoc, a ``-c`` flag, or
``execute_code``. All three trip command-approval gates that stop an
email-driven agent dead. Passing a path does not.

Exactly ONE file is written: .xlsx when openpyxl is importable, otherwise .csv.
"""

import argparse
import csv
import json
import os
import re
import sys

COLUMNS = [
    "Date",
    "Vendor",
    "Description",
    "Subtotal",
    "Sales Tax",
    "Tip",
    "Total",
    "Payment Method",
    "Last 4",
    "Suggested Account",
    "Schedule C Line",
    "Deductible %",
    "Read Confidence",
    "Account Confidence",
    "Needs Review",
    "Review Reason",
    "Source File",
]

MONEY_COLUMNS = {"Subtotal", "Sales Tax", "Tip", "Total"}

DEFAULT_SCHEME = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "..", "reference", "coding_scheme.json"
)


def load_scheme(path):
    return json.load(open(path, encoding="utf-8"))


def scheme_info_lines(path):
    """Return scheme identity and counts derived from the file itself."""
    cfg = load_scheme(path)
    accounts = cfg.get("accounts", [])
    fallbacks = cfg.get("fallback_accounts", [])
    assignable = [a for a in accounts if a.get("auto_assign")]
    not_assignable = [a for a in accounts if not a.get("auto_assign")]

    lines = [
        "scheme_id: %s" % cfg.get("scheme_id", ""),
        "display_name: %s" % cfg.get("display_name", ""),
        "version: %s" % cfg.get("version", ""),
        "source: %s" % cfg.get("source", ""),
        "accounts: %d" % len(accounts),
        "  assignable_from_a_receipt: %d" % len(assignable),
        "  not_assignable: %d" % len(not_assignable),
        "fallback_accounts: %d" % len(fallbacks),
    ]
    for f in fallbacks:
        lines.append("  - %s" % f.get("account", ""))
    lines.append("total_named_accounts: %d" % (len(accounts) + len(fallbacks)))
    return lines


def scheme_info(path, info_file=None):
    """Print scheme counts, and optionally write them to a file.

    The file exists so nothing has to RUN to answer a count question. Across
    three runs the agent improvised execute_code, a `cat | python3 -c`
    pipeline, and a heredoc rather than use the documented command, and each
    improvisation tripped a command-approval gate that emailed the client a
    warning. A plain file read touches no approval system at all.
    """
    lines = scheme_info_lines(path)
    text = "\n".join(lines) + "\n"
    for line in lines:
        print(line)
    if info_file:
        # newline="\n" pins LF on every platform. Without it Python translates
        # to CRLF on Windows, so a committed build artifact would differ
        # depending on which machine generated it.
        tmp = info_file + ".tmp"
        with open(tmp, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(text)
        os.replace(tmp, info_file)
        print("wrote %s" % info_file, file=sys.stderr)
    return 0


def as_int(value):
    """Return an int, or None when the value is blank or unparseable."""
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def as_money(value):
    """Return a float, or None when blank or unparseable."""
    if value is None or value == "":
        return None
    try:
        return round(float(str(value).replace("$", "").replace(",", "").strip()), 2)
    except (TypeError, ValueError):
        return None


REDACTED_REASON = ("This receipt could not be read. No amount and no date were "
                   "recoverable from the image, so every figure has been left "
                   "blank. Check it against the original before entering it.")


def redact_unreadable(row, label, notes, note_text, unreadable):
    """Strip every inferred figure from a row whose source could not be read.

    Money, date and the scheme-derived fields go blank, the description keeps
    only its non-numeric words, and the review reason is REPLACED rather than
    appended to. Replacement matters: on 2026-07-29 the model's own reason read
    "Receipt total is unreadable - shows ??.?? gallons at $3.47/gal with total
    ??.??", quoting a unit price that was never legible on the page. A reason
    written from an unreliable reading is itself unreliable.
    """
    for f in ("Subtotal", "Sales Tax", "Tip", "Total", "Date",
              "Schedule C Line", "Deductible %"):
        row[f] = ""
    desc = scrub_numbers(row.get("Description"))
    row["Description"] = ("%s (amounts unreadable)" % desc if desc
                          else "Contents could not be read")
    row["Suggested Account"] = unreadable
    row["Needs Review"] = "Yes"
    row["Review Reason"] = REDACTED_REASON
    notes.append(note_text)


def scrub_numbers(text):
    """Drop every whitespace token that contains a digit.

    Used when the transcript proves the receipt could not be read. Keeps the
    non-numeric context a bookkeeper needs ("Unleaded gasoline") and removes
    anything that could be mistaken for a figure taken off the page.
    """
    kept = [t for t in str(text or "").split() if not any(c.isdigit() for c in t)]
    out = re.sub(r"\s{2,}", " ", " ".join(kept))
    return out.strip(" -/,;:@.")


def add_reason(row, text):
    existing = str(row.get("Review Reason") or "").strip()
    row["Review Reason"] = ("%s; %s" % (existing, text)) if existing else text


def validate(rows, scheme_path):
    """Force every row into agreement with the scheme. Returns a note list."""
    cfg = load_scheme(scheme_path)
    accounts = {a.get("account"): a for a in cfg.get("accounts", [])}
    fallbacks = {f.get("account"): f for f in cfg.get("fallback_accounts", [])}
    policy = cfg.get("confidence_policy", {})

    floor = policy.get("route_to_fallback_below", 60)
    read_floor = policy.get("read_confidence_floor", 50)
    ambiguous = policy.get("fallback_account_low_confidence", "Ask My Accountant")
    unreadable = policy.get("fallback_account_unreadable", "Uncategorized Expense")
    personal_indicators = [str(k).lower() for k in cfg.get("personal_indicators", [])]

    notes = []
    for n, row in enumerate(rows, start=1):
        label = str(row.get("Vendor") or "row %d" % n)
        acct = str(row.get("Suggested Account") or "").strip()

        # 1. Unreadable receipts cannot be coded at all.
        #
        # This rule used to only re-route the account, and only when the model
        # had not already routed it itself ("if acct != unreadable"). Both
        # halves were wrong.
        #
        # The guard meant the clearest case was the one that got skipped: on
        # 2026-07-29 a fuel receipt came back at read confidence 25 with the
        # account already set to Uncategorized Expense, so this rule saw
        # nothing to do and left "07/17/2026" sitting in the Date field. The
        # model had done the right thing with the account and the wrong thing
        # with the date, and correctly routing itself bought it a free pass.
        #
        # Re-routing alone was also not enough. The row still carried a date
        # and a description built from a reading the model itself scored 25.
        #
        # Low confidence is now an INDEPENDENT reason to strip every figure,
        # and it does not depend on the transcript containing evidence. The
        # "?" rule below only fires if a "?" survived into _raw_text, and that
        # is unreliable in a specific way: fully occluded tokens like "??.??"
        # are preserved, but a token missing a single character from a rigid
        # format gets completed, because completing it is what filling that
        # field looks like. "07/1?/2026" became "07/17/2026" and "3.4?" became
        # "3.47" in the same transcript that faithfully kept two "??.??".
        # On a receipt where every ambiguity is isolated, no "?" is left to
        # find and the "?" rule sees a clean row.
        #
        # Self-reported confidence is imperfect and asymmetric: high confidence
        # proves nothing, but low confidence is the model saying plainly that
        # the figures should not be trusted. There is no reason to argue.
        read_conf = as_int(row.get("Read Confidence"))
        if read_conf is not None and read_conf < read_floor:
            redact_unreadable(
                row, label, notes,
                "%s: read confidence %d is below the floor of %d; money, date and "
                "description blanked and routed to %s"
                % (label, read_conf, read_floor, unreadable),
                unreadable)
            row["Read Confidence"] = 0
            acct = unreadable

        # 2. An account absent from the scheme cannot reach the sheet.
        elif acct not in accounts and acct not in fallbacks:
            notes.append("%s: account %r is not in the scheme, routed to %s"
                         % (label, acct, ambiguous))
            row["Suggested Account"] = ambiguous
            row["Schedule C Line"] = ""
            add_reason(row, "Account not present in the coding scheme")
            acct = ambiguous

        # 3. auto_assign:false accounts are not assignable from a receipt.
        elif acct in accounts and not accounts[acct].get("auto_assign"):
            notes.append("%s: %s is not assignable from a receipt, routed to %s"
                         % (label, acct, ambiguous))
            row["Suggested Account"] = ambiguous
            row["Schedule C Line"] = ""
            add_reason(row, "%s is not assignable from a receipt; needs a human decision" % acct)
            acct = ambiguous

        # 3a. Illegible characters in the transcript mean the numbers were NOT read.
        #
        # On 2026-07-28 a faded fuel receipt reading "TOTAL ??.??" and dated
        # "07/1?/2026" came back as total 77.77 on 7/17/2026 at read
        # confidence 85. Every "?" had been resolved to a "7". Nothing else in
        # this validator could see it, because read confidence is self-reported
        # and an overconfident misread contradicts nothing.
        #
        # So the skill now transcribes VERBATIM into _raw_text before
        # interpreting anything, preserving unresolvable characters as "?".
        # Transcribing what is on the page is a weaker ask than reading it
        # correctly, and it gives us something mechanical to check.
        #
        # A "?" touching a digit means a number on this receipt could not be
        # read. We blank the money and the date rather than keep a value that
        # was inferred, because a fabricated total that looks authoritative is
        # worse than an empty cell that says "look at this one".
        raw = str(row.get("_raw_text") or "")
        if raw:
            if re.search(r"\?\s*\?|\d\s*\?|\?\s*\d|\?\.\d|\d\.\?", raw):
                redact_unreadable(
                    row, label, notes,
                    "%s: transcript contains unreadable characters next to digits; "
                    "money, date and description blanked and routed to %s" % (label, unreadable),
                    unreadable)
                row["Read Confidence"] = 0
                acct = unreadable

        else:
            notes.append("%s: no _raw_text transcript supplied, cannot verify the amounts were read" % label)
            add_reason(row, "No verbatim transcript captured; amounts unverified")
            row["Needs Review"] = "Yes"

        # 3b. Line items that look personal cannot ride a coded business account.
        #
        # This is the one wrong-but-VALID case we can detect from the data
        # already in the row. Office Expense is a real account with
        # auto_assign:true, so checks 1-3 all pass while the coding is still
        # wrong. On 2026-07-28 a Costco receipt with eight grocery and
        # household lines was coded entirely to Office Expense at confidence
        # 85 because the sending email asked for it. The scheme lists
        # "receipt contains both business and personal line items" as an
        # always-flag condition, and that condition must not be waivable by
        # whoever sent the mail.
        #
        # Accounts carrying skip_personal_check are exempt: a restaurant meal
        # is made of food, and flagging it would be a worse bug than the one
        # this catches.
        if acct not in fallbacks and not accounts.get(acct, {}).get("skip_personal_check"):
            items = row.get("_line_items") or []
            if isinstance(items, str):
                items = [items]
            if not items and row.get("Description"):
                items = [str(row.get("Description"))]
            haystack = " ; ".join(str(i).lower() for i in items)
            hits = sorted({k for k in personal_indicators if k and k in haystack})
            if hits:
                shown = ", ".join(hits[:4])
                notes.append("%s: personal-looking line items (%s) cannot sit on %s, routed to %s"
                             % (label, shown, acct, ambiguous))
                row["Suggested Account"] = ambiguous
                row["Schedule C Line"] = ""
                add_reason(row, "Receipt contains personal or household line items (%s); needs a split" % shown)
                acct = ambiguous

        # 3c. Arithmetic must reconcile, and a total with no subtotal is unverified.
        if acct not in fallbacks and acct != unreadable:
            sub, tax, tip, tot = (as_money(row.get(k)) for k in ("Subtotal", "Sales Tax", "Tip", "Total"))
            if sub is not None and tot is not None:
                calc = round(sub + (tax or 0) + (tip or 0), 2)
                if abs(calc - tot) > 0.02:
                    notes.append("%s: %.2f + %.2f + %.2f = %.2f does not match stated total %.2f"
                                 % (label, sub, tax or 0, tip or 0, calc, tot))
                    add_reason(row, "Amounts do not reconcile (%.2f vs stated %.2f)" % (calc, tot))
                    row["Needs Review"] = "Yes"
            elif tot is not None and sub is None:
                read_c = as_int(row.get("Read Confidence"))
                if read_c is None or read_c < 95:
                    notes.append("%s: total %.2f has no subtotal to check it against at read confidence %s, "
                                 "routed to %s" % (label, tot, read_c, ambiguous))
                    row["Suggested Account"] = ambiguous
                    row["Schedule C Line"] = ""
                    add_reason(row, "Total could not be verified against a subtotal")
                    acct = ambiguous

        # 3d. Some accounts are never clean on the strength of a receipt alone.
        #
        # IRC 274(d) requires business purpose and attendees for meals. A
        # receipt shows what was bought, never why or with whom, so no meal
        # can be substantiated from the image. The scheme said so in prose
        # ("Business purpose and attendees must be recorded") and the model
        # honoured it on 2026-07-28 and dropped it on 2026-07-29, returning
        # Needs Review: No on a Panera receipt. Prose has now failed on this
        # rule twice, so it moves into the validator.
        #
        # Driven by requires_substantiation in the scheme rather than an
        # account name, so adding Travel later is a one-line data change.
        if acct in accounts and accounts[acct].get("requires_substantiation"):
            if str(row.get("Needs Review") or "").strip().lower() not in ("yes", "true", "y", "1"):
                notes.append("%s: %s always needs substantiation, forced Needs Review"
                             % (label, acct))
                row["Needs Review"] = "Yes"
            add_reason(row, "Business purpose and attendees are not on the receipt "
                            "and must be recorded before this is deductible")

        # 4. A fallback row cannot claim high account confidence, and is always flagged.
        if acct in fallbacks:
            acct_conf = as_int(row.get("Account Confidence"))
            if acct_conf is None or acct_conf >= floor:
                corrected = floor - 1
                if acct_conf is not None:
                    notes.append("%s: account confidence %d contradicts routing to %s, set to %d"
                                 % (label, acct_conf, acct, corrected))
                row["Account Confidence"] = corrected
            if str(row.get("Needs Review") or "").strip().lower() not in ("yes", "true", "y", "1"):
                notes.append("%s: routed to %s, forced Needs Review" % (label, acct))
                row["Needs Review"] = "Yes"
            row["Deductible %"] = ""

    return notes


def normalize(rows):
    out = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("each row must be a JSON object")
        out.append({c: ("" if row.get(c) is None else row.get(c)) for c in COLUMNS})
    return out


def build_summary(rows):
    """Compose the reply body from VALIDATED rows.

    Why this is not left to the model
    ---------------------------------
    On 2026-07-29 the sheet was correct and the covering email was not. The
    Shell row had its date blanked by the validator, and the email still
    opened "Shell (07/17/2026) - Receipt total is unreadable", quoting a date
    that is not on the receipt and not on the sheet. The email had been
    written from the model's own pre-validation row.

    That is the third time the same invented figure relocated rather than
    disappeared: first the Total column, then the Description, then the
    covering prose. Each fix worked and each time the guess moved to the
    nearest surface the validator did not cover. Prose is the last such
    surface, so it stops being prose.

    Everything below is derived from `rows` AFTER validation. A value the
    validator blanked cannot reappear here, because there is nothing to read
    it from.
    """
    def flagged(r):
        return str(r.get("Needs Review") or "").strip().lower() in ("yes", "true", "y", "1")

    total = len(rows)
    review = [r for r in rows if flagged(r)]
    clean = total - len(review)

    noun = "receipt" if total == 1 else "receipts"
    lines = ["I read %d %s. %d coded cleanly, %d need review."
             % (total, noun, clean, len(review))]

    if review:
        lines.append("")
        lines.append("Flagged for review:")
        for r in review:
            vendor = str(r.get("Vendor") or "").strip() or "Unknown vendor"
            date = str(r.get("Date") or "").strip()
            amount = str(r.get("Total") or "").strip()
            reason = " ".join(str(r.get("Review Reason") or "").split()) or "Flagged for review."

            # Blank money or date is not an omission to paper over. It is the
            # validator reporting the figure was never legible, and the reply
            # has to say so rather than quietly leaving it out.
            if date and amount:
                ident = "%s, %s, %s" % (vendor, date, amount)
            elif date:
                ident = "%s, %s, amount could not be read" % (vendor, date)
            elif amount:
                ident = "%s, %s, date could not be read" % (vendor, amount)
            else:
                ident = "%s, date and amount could not be read" % vendor
            lines.append("- %s: %s" % (ident, reason))

    return "\n".join(lines) + "\n"


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    return path


def write_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Approval Sheet"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    flag_fill = PatternFill("solid", fgColor="FFF2CC")

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    review_col = COLUMNS.index("Needs Review")
    for row in rows:
        ws.append([row[c] for c in COLUMNS])
        flagged = str(row[COLUMNS[review_col]]).strip().lower() in ("yes", "true", "y", "1")
        if flagged:
            for cell in ws[ws.max_row]:
                cell.fill = flag_fill

    for i, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        width = max(len(col), *(len(str(r[col])) for r in rows)) if rows else len(col)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 45)
        if col in MONEY_COLUMNS:
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0.00"

    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def read_rows(args):
    if args.rows_file:
        raw = open(args.rows_file, encoding="utf-8").read().strip()
        origin = args.rows_file
    else:
        raw = sys.stdin.read().strip()
        origin = "stdin"
    if not raw:
        raise ValueError("no JSON received from %s" % origin)
    rows = json.loads(raw)
    if not isinstance(rows, list):
        raise ValueError("expected a JSON array of row objects from %s" % origin)
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scheme-info", action="store_true", help="print scheme identity and counts")
    ap.add_argument("--scheme", default=DEFAULT_SCHEME, help="path to coding_scheme.json")
    ap.add_argument("--rows-file", help="path to a JSON array of row objects")
    ap.add_argument("--out", help="output path WITHOUT extension")
    ap.add_argument("--info-file", help="also write --scheme-info output to this path")
    ap.add_argument("--summary-file",
                    help="write the reply body, derived from the validated rows, to this path")
    ap.add_argument("--no-validate", action="store_true", help="skip the consistency pass (testing only)")
    args = ap.parse_args()

    if args.scheme_info:
        try:
            return scheme_info(args.scheme, args.info_file)
        except Exception as exc:
            print("ERROR: could not read scheme: %s" % exc, file=sys.stderr)
            return 2

    if not args.out:
        print("ERROR: --out is required when rendering a sheet", file=sys.stderr)
        return 2

    try:
        raw_rows = read_rows(args)
    except (ValueError, json.JSONDecodeError) as exc:
        print("ERROR: %s" % exc, file=sys.stderr)
        return 2
    except OSError as exc:
        print("ERROR: could not read rows file: %s" % exc, file=sys.stderr)
        return 2

    # Validation runs on the RAW rows so helper keys like _line_items are still
    # present; normalize() drops everything outside COLUMNS.
    if not args.no_validate:
        try:
            notes = validate(raw_rows, args.scheme)
        except Exception as exc:
            print("ERROR: validation could not run: %s" % exc, file=sys.stderr)
            return 2
        if notes:
            print("VALIDATION: %d row(s) corrected to match the scheme:" % len(notes), file=sys.stderr)
            for note in notes:
                print("  - %s" % note, file=sys.stderr)

    try:
        rows = normalize(raw_rows)
    except ValueError as exc:
        print("ERROR: %s" % exc, file=sys.stderr)
        return 2

    written = write_xlsx(rows, args.out + ".xlsx")
    if not written:
        written = write_csv(rows, args.out + ".csv")
        print("NOTE: openpyxl unavailable, wrote CSV instead", file=sys.stderr)

    if args.summary_file:
        try:
            with open(args.summary_file, "w", encoding="utf-8") as fh:
                fh.write(build_summary(rows))
        except OSError as exc:
            # The sheet is the deliverable; a summary failure must not lose it.
            print("WARNING: could not write summary file: %s" % exc, file=sys.stderr)
        else:
            # Labelled, and printed BEFORE the sheet path, so the last line of
            # stdout is still the single path to attach exactly as before.
            print("SUMMARY: %s" % args.summary_file)

    print(written)
    return 0


if __name__ == "__main__":
    sys.exit(main())
